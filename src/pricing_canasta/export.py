import json
from datetime import UTC, date, datetime
from itertools import islice
from pathlib import Path
from uuid import UUID

import duckdb
import pandas as pd
import plotly.graph_objects as go
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from plotly.subplots import make_subplots

from .config import (
    DATABASE_PATH,
    EXPORT_DIR,
    OUTPUT_DIR,
    PORTFOLIO_DATA_DIR,
    SITE_DIR,
    ensure_directories,
)
from .validate import ValidationResult, validate

SMALL_EXPORTS = {
    "quality_daily": "SELECT * FROM mart_quality_daily ORDER BY snapshot_date",
    "quality_checks": "SELECT * FROM quality_checks ORDER BY severity, test_name",
    "source_health": "SELECT * FROM mart_snapshot_health ORDER BY snapshot_date",
    "banner_health": "SELECT * FROM mart_banner_day_health ORDER BY snapshot_date, banner_label",
    "basket_definition": """
        SELECT c.*, p.producto_descripcion, p.marca, p.cantidad_base, p.unidad_base
        FROM dim_basket_component AS c
        LEFT JOIN mart_basket_candidate_review AS p
            ON c.build_id = p.build_id AND c.basket_version = p.basket_version
            AND c.category_id = p.category_id AND c.gtin14 = p.gtin14
            AND p.selected_component
        ORDER BY c.component_order
    """,
    "basket_candidate_review": """
        SELECT * FROM mart_basket_candidate_review
        WHERE selected_component OR (semantic_match AND presentation_match)
        ORDER BY category_id, selected_component DESC, gtin14
    """,
    "index_daily": """
        SELECT i.*, s.banner_label
        FROM mart_banner_index_daily AS i
        LEFT JOIN dim_banner_scope AS s
            USING (build_id, scope_version, id_comercio, id_bandera)
        ORDER BY i.snapshot_date, i.price_index
    """,
    "index_7d": """
        SELECT i.*, s.banner_label
        FROM mart_banner_index_7d AS i
        LEFT JOIN dim_banner_scope AS s
            USING (build_id, scope_version, id_comercio, id_bandera)
        ORDER BY i.geometric_mean_index
    """,
    "index_sensitivity": """
        SELECT i.*, s.banner_label,
               abs(i.raw_index - i.exclude_critical_index) AS sensitivity_delta
        FROM mart_banner_index_sensitivity_daily AS i
        LEFT JOIN dim_banner_scope AS s
            USING (build_id, scope_version, id_comercio, id_bandera)
        ORDER BY i.snapshot_date, i.raw_index
    """,
    "index_drivers": """
        SELECT d.*, s.banner_label, p.producto_descripcion, p.marca
        FROM mart_banner_index_driver_daily AS d
        LEFT JOIN dim_banner_scope AS s
            USING (build_id, scope_version, id_comercio, id_bandera)
        LEFT JOIN dim_product_asof AS p USING (build_id, gtin14)
        WHERE d.driver_rank <= 15
        ORDER BY d.snapshot_date, s.banner_label, d.driver_rank
    """,
    "banner_overlap": "SELECT * FROM mart_banner_pair_overlap_daily ORDER BY snapshot_date, common_gtins DESC",
    "dispersion_entity": """
        SELECT d.*, coalesce(s.banner_label, b.comercio_bandera_nombre) AS banner_label
        FROM mart_dispersion_entity_daily AS d
        LEFT JOIN dim_banner_daily AS b
            ON d.build_id = b.build_id AND d.snapshot_date = b.snapshot_date
            AND d.id_comercio = b.id_comercio AND d.id_bandera = b.id_bandera
        LEFT JOIN dim_banner_scope AS s
            ON d.build_id = s.build_id AND d.id_comercio = s.id_comercio
            AND d.id_bandera = s.id_bandera
        ORDER BY d.snapshot_date, d.dispersion_level, d.median_dispersion_clean DESC
    """,
    "dispersion_product_sample": """
        SELECT d.*, p.producto_descripcion, p.marca
        FROM mart_price_dispersion_daily AS d
        LEFT JOIN dim_product_asof AS p USING (build_id, gtin14)
        WHERE d.coverage_status IN ('PUBLISHABLE', 'DIRECTIONAL')
        ORDER BY d.snapshot_date DESC, d.dispersion_clean DESC NULLS LAST
        LIMIT 10000
    """,
    "basket_banner": """
        SELECT b.*, coalesce(sc.banner_label, s.comercio_bandera_nombre) AS banner_label
        FROM mart_basket_banner_daily AS b
        LEFT JOIN dim_banner_daily AS s
            ON b.build_id = s.build_id AND b.snapshot_date = s.snapshot_date
            AND b.id_comercio = s.id_comercio AND b.id_bandera = s.id_bandera
        LEFT JOIN dim_banner_scope AS sc
            ON b.build_id = sc.build_id AND b.id_comercio = sc.id_comercio
            AND b.id_bandera = sc.id_bandera
        ORDER BY b.snapshot_date, b.branch_median_cost
    """,
    "basket_province": "SELECT * FROM mart_basket_province_daily ORDER BY snapshot_date, branch_median_cost",
    "basket_banner_province": """
        SELECT b.*, coalesce(sc.banner_label, s.comercio_bandera_nombre) AS banner_label
        FROM mart_basket_banner_province_daily AS b
        LEFT JOIN dim_banner_daily AS s
            ON b.build_id = s.build_id AND b.snapshot_date = s.snapshot_date
            AND b.id_comercio = s.id_comercio AND b.id_bandera = s.id_bandera
        LEFT JOIN dim_banner_scope AS sc
            ON b.build_id = sc.build_id AND b.id_comercio = sc.id_comercio
            AND b.id_bandera = sc.id_bandera
        ORDER BY b.snapshot_date, b.provincia_codigo, b.branch_median_cost
    """,
    "basket_national": "SELECT * FROM mart_basket_national_daily ORDER BY snapshot_date",
    "basket_savings": "SELECT * FROM mart_basket_savings_daily ORDER BY snapshot_date, comparison_level",
    "price_anomalies_sample": """
        SELECT a.*, p.producto_descripcion, p.marca
        FROM quality_price_anomaly_daily AS a
        LEFT JOIN dim_product_asof AS p USING (build_id, gtin14)
        ORDER BY CASE a.severity WHEN 'CRITICAL' THEN 1 ELSE 2 END,
                 a.snapshot_date DESC,
                 abs(ln(a.price_ratio)) DESC
        LIMIT 10000
    """,
}

DASHBOARD_CONFIG = {
    "responsive": True,
    "displayModeBar": True,
    "displaylogo": False,
    "scrollZoom": False,
    "doubleClick": "reset+autosize",
    "locale": "es",
    "locales": {
        "es": {
            "dictionary": {
                "Autoscale": "Ajustar escala",
                "Download plot as a PNG": "Descargar gráfico como PNG",
                "Pan": "Desplazar",
                "Reset axes": "Restablecer ejes",
                "Zoom": "Ampliar o reducir",
                "Zoom in": "Ampliar",
                "Zoom out": "Reducir",
            },
            "format": {
                "shortDays": ["Dom", "Lun", "Mar", "Mié", "Jue", "Vie", "Sáb"],
                "shortMonths": [
                    "Ene",
                    "Feb",
                    "Mar",
                    "Abr",
                    "May",
                    "Jun",
                    "Jul",
                    "Ago",
                    "Sep",
                    "Oct",
                    "Nov",
                    "Dic",
                ],
                "date": "%d/%m/%Y",
                "decimal": ",",
                "thousands": ".",
            },
        }
    },
    "modeBarButtonsToAdd": ["zoom2d", "pan2d", "autoScale2d", "resetScale2d"],
    "modeBarButtonsToRemove": ["select2d", "lasso2d"],
}

LEGACY_EXPORT_FILES = {
    "basket_summary",
    "dispersion_product",
    "ingestion_log",
    "price_anomalies",
    "price_dispersion",
    "price_index_banner",
    "product_candidates",
}

LARGE_PARQUET_EXPORTS = {
    "price_product_daily": "SELECT * FROM mart_price_product_daily",
    "basket_store_daily": "SELECT * FROM mart_basket_store_daily",
    "basket_component_store_daily": "SELECT * FROM mart_basket_component_store_daily",
    "product_dimension": "SELECT * FROM dim_product_asof",
    "store_dimension_daily": "SELECT * FROM dim_store_daily",
}


def _sql_literal(value: object) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def _export_tables(connection: duckdb.DuckDBPyConnection) -> dict[str, pd.DataFrame]:
    for directory in (EXPORT_DIR, PORTFOLIO_DATA_DIR):
        for name in LEGACY_EXPORT_FILES:
            (directory / f"{name}.csv").unlink(missing_ok=True)
            (directory / f"{name}.parquet").unlink(missing_ok=True)

    frames = {}
    for name, query in SMALL_EXPORTS.items():
        frame = connection.execute(query).fetchdf()
        frames[name] = frame
        for directory in (EXPORT_DIR, PORTFOLIO_DATA_DIR):
            frame.to_csv(directory / f"{name}.csv", index=False, encoding="utf-8-sig")
        frame.to_parquet(EXPORT_DIR / f"{name}.parquet", index=False)
        print(f"[export] {name}: {len(frame):,} filas")

    for name, query in LARGE_PARQUET_EXPORTS.items():
        destination = EXPORT_DIR / f"{name}.parquet"
        destination.unlink(missing_ok=True)
        connection.execute(
            f"COPY ({query}) TO {_sql_literal(destination.resolve())} "
            "(FORMAT PARQUET, COMPRESSION ZSTD)"
        )
        print(f"[export] {name}: Parquet generado")
    return frames


def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="17324D")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.columns, start=1):
        width = max(len(str(cell.value or "")) for cell in islice(column, 500)) + 2
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width, 10), 45)


def _excel_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    if isinstance(value, UUID):
        return str(value)
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def _write_excel(frames: dict[str, pd.DataFrame], validation: ValidationResult) -> Path:
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    rows = [
        ("Proyecto", "Pricing y canasta de supermercado v2"),
        ("Versión", validation.build_id),
        ("Fecha de corte", str(validation.as_of_date)),
        ("Confianza", validation.confidence),
        ("Índice", "Panel común de GTIN; igual peso por producto y cadena."),
        ("Dispersión", "P90-P10 relativo; valores originales y sin anomalías críticas."),
        ("Canasta", "Ocho GTIN fijos; solo canastas completas; no se imputan faltantes."),
        ("Nacional", "Red de sucursales observada, no ponderada por poblacion."),
    ]
    readme.append(["Campo", "Detalle"])
    for row in rows:
        readme.append(row)
    _style_sheet(readme)

    sheet_order = [
        "quality_daily",
        "quality_checks",
        "source_health",
        "basket_definition",
        "index_daily",
        "index_7d",
        "index_sensitivity",
        "index_drivers",
        "dispersion_entity",
        "dispersion_product_sample",
        "basket_banner",
        "basket_province",
        "basket_banner_province",
        "basket_national",
        "basket_savings",
        "price_anomalies_sample",
    ]
    for name in sheet_order:
        frame = frames[name]
        sheet = workbook.create_sheet(name[:31])
        sheet.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            sheet.append([_excel_value(value) for value in row])
        _style_sheet(sheet)

    path = OUTPUT_DIR / "pricing_canasta_supermercado.xlsx"
    workbook.save(path)
    return path


def _latest_publishable(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty:
        return frame
    publishable = frame
    if "coverage_status" in frame.columns:
        publishable = frame[frame["coverage_status"] == "PUBLISHABLE"]
    if publishable.empty:
        return frame.iloc[0:0]
    latest = publishable["snapshot_date"].max()
    return publishable[publishable["snapshot_date"] == latest]


def _latest_publishable_level(frame: pd.DataFrame, level: str) -> pd.DataFrame:
    return _latest_publishable(frame[frame["dispersion_level"] == level])


def _date_label(frame: pd.DataFrame) -> str:
    if frame.empty or "snapshot_date" not in frame:
        return "sin dato publicable"
    return pd.Timestamp(frame["snapshot_date"].max()).strftime("%d/%m/%Y")


def _dispersion_panel_text(
    label: str, frame: pd.DataFrame, as_of_date: date
) -> tuple[str, str]:
    if frame.empty:
        title = f"{label} | sin dato publicable"
        notice = f"{label}: sin dato publicable en la ventana observada."
        return title, notice

    effective_date = pd.Timestamp(frame["snapshot_date"].max())
    effective_label = effective_date.strftime("%d/%m/%Y")
    title = f"{label} | dato efectivo {effective_label}"
    notice = f"{label}: dato efectivo {effective_label}."
    if effective_date.date() < as_of_date:
        global_label = as_of_date.strftime("%d/%m/%Y")
        warning = f"El corte global del {global_label} no tuvo cobertura publicable para este panel."
        title += (
            f"<br><span style='color:#f3ce62'>Corte {global_label}: "
            "sin cobertura publicable</span>"
        )
        notice += f" {warning}"
    return title, notice


def _enhance_dashboard_html(
    path: Path, validation: ValidationResult, panel_notices: tuple[str, ...]
) -> None:
    html = path.read_text(encoding="utf-8")
    html = html.replace("<html>", '<html lang="es-AR">', 1)
    html = html.replace(
        "<head>",
        (
            '<head><meta name="viewport" '
            'content="width=device-width, initial-scale=1">'
            "<title>Pricing y canasta de supermercado</title>"
            "<style>@import url('https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:"
            "wght@400;500;600&family=Space+Grotesk:wght@400;500;600;700&display=swap');"
            ":root{--bg:#0b1111;--bg-deep:#080d0d;--surface:#171d1e;"
            "--surface-raised:#1b2223;--surface-soft:#111819;--text:#edf1ef;"
            "--muted:#a5aeaa;--dim:#89938f;--mint:#9ef6e5;"
            "--mint-bright:#58e4d0;--amber:#f3ce62;"
            "--border:rgba(158,246,229,.11);--border-strong:rgba(158,246,229,.25);}"
            "*{box-sizing:border-box;}html{background:var(--bg-deep);}"
            "body{margin:0;color:var(--text);background-color:var(--bg);"
            "background-image:linear-gradient(var(--border) 1px,transparent 1px),"
            "linear-gradient(90deg,var(--border) 1px,transparent 1px);"
            "background-size:48px 48px;font-family:'Space Grotesk',Arial,sans-serif;}"
            ".site-header,.dashboard-shell{width:min(1480px,calc(100% - 40px));margin:0 auto;}"
            ".site-header{padding:42px 0 24px;border-bottom:1px solid var(--border-strong);}"
            ".eyebrow,.site-meta,.view-help,.back-link,.reset-view{"
            "font-family:'IBM Plex Mono',Consolas,monospace;}"
            ".eyebrow{margin:0 0 12px;color:var(--mint-bright);font-size:.75rem;"
            "letter-spacing:.12em;text-transform:uppercase;}"
            "h1{max-width:900px;margin:0;font-size:clamp(2rem,5vw,4.6rem);"
            "font-weight:600;line-height:.95;letter-spacing:-.045em;}"
            ".site-description{max-width:780px;margin:22px 0 18px;color:var(--muted);"
            "overflow-wrap:anywhere;"
            "font-size:clamp(1rem,2vw,1.2rem);line-height:1.55;}"
            ".site-meta{display:flex;flex-wrap:wrap;gap:8px 24px;color:var(--dim);font-size:.78rem;}"
            ".site-meta span{min-width:0;overflow-wrap:anywhere;}"
            ".back-link{display:inline-block;margin-top:24px;color:var(--mint);"
            "text-underline-offset:4px;}"
            ".back-link:hover,.back-link:focus-visible{color:var(--text);}"
            ".dashboard-shell{padding:24px 0 48px;}"
            ".panel-notices{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));"
            "border-top:1px solid var(--border);border-left:1px solid var(--border);"
            "background:var(--surface-soft);}"
            ".panel-notices p{margin:0;padding:14px 16px;border-right:1px solid var(--border);"
            "border-bottom:1px solid var(--border);color:var(--amber);font-size:.88rem;line-height:1.45;}"
            ".dashboard-toolbar{display:flex;align-items:center;justify-content:space-between;"
            "gap:16px;padding:14px 0;}"
            ".view-help{margin:0;color:var(--dim);font-size:.75rem;line-height:1.45;}"
            ".reset-view{min-height:44px;padding:10px 14px;border:1px solid var(--border-strong);"
            "border-radius:0;color:var(--bg-deep);background:var(--mint);font-weight:700;cursor:pointer;}"
            ".reset-view:hover,.reset-view:focus-visible{background:var(--mint-bright);outline:none;}"
            ".plotly-graph-div{width:100%!important;max-width:100%;"
            "border:1px solid var(--border);background:var(--surface-soft);box-shadow:none;}"
            ".modebar{opacity:1!important;}"
            ".sr-only{position:absolute;width:1px;height:1px;padding:0;margin:-1px;"
            "overflow:hidden;clip:rect(0,0,0,0);white-space:nowrap;border:0;}"
            "@media(max-width:700px){.site-header,.dashboard-shell{width:min(100% - 24px,1480px);}"
            ".site-header{padding-top:28px}.panel-notices{grid-template-columns:1fr;}"
            ".dashboard-toolbar{align-items:stretch;flex-direction:column;}"
            ".reset-view{width:100%;}.view-help{order:2;}}</style>"
        ),
        1,
    )
    notices = "".join(f"<p>{notice}</p>" for notice in panel_notices)
    cut_label = validation.as_of_date.strftime("%d/%m/%Y")
    html = html.replace(
        "<body>",
        (
            '<body><header class="site-header"><p class="eyebrow">Datos abiertos · Argentina</p>'
            "<h1>Pricing y canasta de supermercado</h1>"
            '<p class="site-description">Nivel de precios, dispersión y costo de una canasta '
            "comparable, con cobertura y fechas efectivas visibles para interpretar cada panel.</p>"
            f'<div class="site-meta"><span>Corte global: {cut_label}</span>'
            "<span>Fuente: Precios Claros · Secretaría de Industria y Comercio</span>"
            f"<span>Confianza del corte: {validation.confidence}</span></div>"
            '<a class="back-link" href="https://alexanderhuth98.github.io/#proyectos">'
            "← Volver a proyectos</a></header>"
            '<main class="dashboard-shell"><p class="sr-only">Tablero de nivel de precios, '
            "canasta, dispersión, sensibilidad y cobertura. El índice identifica el último día "
            "comparable. Cada panel informa su fecha efectiva; los datos suprimidos no se incluyen "
            "en clasificaciones.</p>"
            f'<section class="panel-notices" aria-label="Vigencia de los paneles">{notices}</section>'
            '<div class="dashboard-toolbar"><p class="view-help">Usá la barra del gráfico para '
            "ampliar, desplazar o ajustar los ejes.</p>"
            '<button class="reset-view" type="button">Restablecer vista</button></div>'
        ),
        1,
    )
    html = html.replace(
        '<div id="',
        '<div role="img" aria-label="Dashboard interactivo de pricing y canasta" id="',
        1,
    )
    html = html.replace(
        "</body>",
        """<script>
(() => {
    const button = document.querySelector('.reset-view');
    const graph = document.querySelector('.plotly-graph-div');
    if (!button || !graph) return;
    button.addEventListener('click', () => {
        const update = {autosize: true};
        Object.keys(graph.layout || {}).forEach((key) => {
            if (/^[xy]axis\\d*$/.test(key)) update[`${key}.autorange`] = true;
        });
        Plotly.relayout(graph, update);
    });
})();
</script></main></body>""",
        1,
    )
    html = "\n".join(line.rstrip() for line in html.splitlines()) + "\n"
    path.write_text(html, encoding="utf-8", newline="\n")


def _write_dashboard(
    frames: dict[str, pd.DataFrame], validation: ValidationResult, mobile: bool = False
) -> Path:
    index = frames["index_daily"]
    index = index[index["publishable"]]
    index_latest = (
        index[index["snapshot_date"] == index["snapshot_date"].max()] if not index.empty else index
    )
    basket_banner = _latest_publishable(frames["basket_banner"])
    dispersion = frames["dispersion_entity"]
    chain_disp = _latest_publishable_level(dispersion, "BANNER").sort_values(
        "median_dispersion_clean"
    )
    province_disp = _latest_publishable_level(dispersion, "PROVINCE").sort_values(
        "median_dispersion_clean"
    )
    health = frames["source_health"]
    sensitivity = frames["index_sensitivity"]
    sensitivity_latest = (
        sensitivity[sensitivity["snapshot_date"] == index_latest["snapshot_date"].max()]
        if not index_latest.empty
        else sensitivity.iloc[0:0]
    )

    index_date = _date_label(index_latest)
    basket_date = _date_label(basket_banner)
    chain_title, chain_notice = _dispersion_panel_text(
        "Dispersión por cadena", chain_disp, validation.as_of_date
    )
    province_title, province_notice = _dispersion_panel_text(
        "Dispersión por provincia", province_disp, validation.as_of_date
    )
    window_start = pd.Timestamp(health["snapshot_date"].min()).strftime("%d/%m")
    window_end = pd.Timestamp(health["snapshot_date"].max()).strftime("%d/%m/%Y")

    subplot_titles = (
        f"GTIN comunes | \u00faltimo d\u00eda comparable {index_date}",
        f"Salud cadena-día | semana {window_start}-{window_end}",
        f"\u00cdndice de precios | \u00faltimo d\u00eda comparable {index_date}",
        f"Canasta publicable por cadena | {basket_date}",
        chain_title,
        province_title,
        f"Sensibilidad a extremos | {index_date}",
        f"Tiendas reportantes vs. mediana 7d | {window_start}-{window_end}",
    )
    if mobile:
        subplot_titles = (
            f"GTIN comunes | {index_date}",
            f"Salud semanal | {window_start}-{window_end}",
            f"\u00cdndice | {index_date}",
            f"Canasta por cadena | {basket_date}",
            chain_title,
            province_title,
            f"Sensibilidad | {index_date}",
            f"Tiendas vs. mediana 7d | {window_start}-{window_end}",
        )
        positions = [(row, 1) for row in range(1, 9)]
        figure = make_subplots(
            rows=8,
            cols=1,
            specs=[[{"type": "indicator"}], [{"type": "indicator"}]] + [[{"type": "xy"}]] * 6,
            subplot_titles=subplot_titles,
            vertical_spacing=0.055,
        )
        height = 2500
    else:
        positions = [(1, 1), (1, 2), (2, 1), (2, 2), (3, 1), (3, 2), (4, 1), (4, 2)]
        figure = make_subplots(
            rows=4,
            cols=2,
            specs=[
                [{"type": "indicator"}, {"type": "indicator"}],
                [{"type": "xy"}, {"type": "xy"}],
                [{"type": "xy"}, {"type": "xy"}],
                [{"type": "xy"}, {"type": "xy"}],
            ],
            subplot_titles=subplot_titles,
            vertical_spacing=0.10,
            horizontal_spacing=0.12,
        )
        height = 1200
    common_gtins = index_latest["common_gtins"].min() if not index_latest.empty else 0
    healthy_ratio = frames["banner_health"]["source_healthy"].mean()
    figure.add_trace(
        go.Indicator(
            mode="number",
            value=float(common_gtins),
            number={"font": {"size": 54, "color": "#9ef6e5"}} if mobile else {
                "font": {"color": "#9ef6e5"}
            },
        ),
        row=positions[0][0],
        col=positions[0][1],
    )
    figure.add_trace(
        go.Indicator(
            mode="number",
            value=float(100 * healthy_ratio),
            number={
                "suffix": "%",
                "valueformat": ".1f",
                "font": {"size": 54 if mobile else 42, "color": "#9ef6e5"},
            },
        ),
        row=positions[1][0],
        col=positions[1][1],
    )

    for banner, group in index.groupby("banner_label", sort=True):
        figure.add_trace(
            go.Scatter(
                x=group["snapshot_date"],
                y=group["price_index"],
                mode="lines+markers",
                name=banner,
                customdata=group[["common_gtins"]],
                hovertemplate=(
                    "%{fullData.name}<br>Fecha: %{x|%d/%m/%Y}"
                    "<br>\u00cdndice: %{y:.2f}<br>GTIN comunes: %{customdata[0]:,.0f}"
                    "<extra></extra>"
                ),
            ),
            row=positions[2][0],
            col=positions[2][1],
        )
    figure.add_shape(
        type="line",
        x0=0,
        x1=1,
        y0=100,
        y1=100,
        xref="x domain",
        yref="y",
        line={"dash": "dash", "color": "#89938f"},
    )

    figure.add_trace(
        go.Bar(
            x=basket_banner["branch_median_cost"],
            y=basket_banner["banner_label"],
            orientation="h",
            marker_color="#f3ce62",
            name="Canasta",
            customdata=basket_banner[["complete_stores", "completion_rate"]],
            hovertemplate="%{y}<br>$%{x:,.0f}<br>Sucursales: %{customdata[0]}<br>Cobertura: %{customdata[1]:.1%}<extra></extra>",
        ),
        row=positions[3][0],
        col=positions[3][1],
    )

    figure.add_trace(
        go.Bar(
            x=chain_disp["median_dispersion_clean"],
            y=chain_disp["banner_label"],
            orientation="h",
            name="Cadena",
            marker_color="#58e4d0",
        ),
        row=positions[4][0],
        col=positions[4][1],
    )
    figure.add_trace(
        go.Bar(
            x=province_disp["median_dispersion_clean"],
            y=province_disp["provincia_codigo"],
            orientation="h",
            name="Provincia",
            marker_color="#9ef6e5",
        ),
        row=positions[5][0],
        col=positions[5][1],
    )
    empty_panels = (
        (chain_disp, positions[4], "x3 domain", "y3 domain"),
        (province_disp, positions[5], "x4 domain", "y4 domain"),
    )
    for frame, position, xref, yref in empty_panels:
        if frame.empty:
            figure.add_annotation(
                text="Sin cobertura publicable en la ventana observada",
                x=0.5,
                y=0.5,
                xref=xref,
                yref=yref,
                showarrow=False,
                font={"size": 14, "color": "#a5aeaa"},
            )
            figure.update_xaxes(visible=False, row=position[0], col=position[1])
            figure.update_yaxes(visible=False, row=position[0], col=position[1])

    for banner, group in sensitivity_latest.groupby("banner_label", sort=True):
        figure.add_trace(
            go.Scatter(
                x=group["raw_index"],
                y=group["exclude_critical_index"],
                mode="markers" if mobile else "markers+text",
                text=None if mobile else [banner] * len(group),
                textposition="top center",
                marker={"size": 11, "color": "#f3ce62"},
                name=f"Sensibilidad | {banner}",
                showlegend=mobile,
                hovertemplate=(
                    f"{banner}<br>\u00cdndice original: %{{x:.2f}}"
                    "<br>Sin cr\u00edticos: %{y:.2f}<extra></extra>"
                ),
            ),
            row=positions[6][0],
            col=positions[6][1],
        )
    if not sensitivity_latest.empty:
        sensitivity_min = float(
            sensitivity_latest[["raw_index", "exclude_critical_index"]].min().min()
        )
        sensitivity_max = float(
            sensitivity_latest[["raw_index", "exclude_critical_index"]].max().max()
        )
        figure.add_trace(
            go.Scatter(
                x=[sensitivity_min, sensitivity_max],
                y=[sensitivity_min, sensitivity_max],
                mode="lines",
                line={"dash": "dot", "color": "#89938f"},
                hoverinfo="skip",
                showlegend=False,
            ),
            row=positions[6][0],
            col=positions[6][1],
        )
    figure.add_trace(
        go.Bar(
            x=health["snapshot_date"],
            y=health["store_coverage_ratio"],
            marker_color=["#58e4d0" if value else "#f3ce62" for value in health["source_healthy"]],
            customdata=pd.DataFrame(
                {
                    "reporting_stores": health["reporting_stores"],
                    "source_status": health["source_healthy"].map({True: "Sí", False: "No"}),
                }
            ),
            hovertemplate=(
                "Fecha: %{x|%d/%m/%Y}<br>Tiendas vs. mediana: %{y:.1%}"
                "<br>Tiendas reportantes: %{customdata[0]:,.0f}"
                "<br>Fuente saludable: %{customdata[1]}<extra></extra>"
            ),
            name="Tiendas vs. mediana 7d",
        ),
        row=positions[7][0],
        col=positions[7][1],
    )

    figure.update_layout(
        template="plotly_dark",
        paper_bgcolor="#0b1111",
        plot_bgcolor="#111819",
        colorway=["#9ef6e5", "#58e4d0", "#f3ce62", "#a5aeaa", "#89938f"],
        font={
            "family": "Space Grotesk, Arial, sans-serif",
            "color": "#edf1ef",
            "size": 11 if mobile else 12,
        },
        height=height,
        width=None,
        legend={
            "orientation": "h",
            "y": -0.04,
            "font": {"size": 9 if mobile else 12},
            "bgcolor": "rgba(11,17,17,.88)",
        },
        margin={
            "l": 45 if mobile else 80,
            "r": 20 if mobile else 40,
            "t": 80 if mobile else 90,
            "b": 80 if mobile else 100,
        },
        autosize=True,
        hoverlabel={
            "bgcolor": "#1b2223",
            "bordercolor": "rgba(158,246,229,.25)",
            "font": {"family": "IBM Plex Mono, Consolas, monospace", "color": "#edf1ef"},
        },
    )
    if mobile:
        figure.update_annotations(font={"size": 11, "color": "#edf1ef"})
        figure.update_xaxes(tickfont={"size": 9}, title_font={"size": 10})
        figure.update_yaxes(tickfont={"size": 9}, title_font={"size": 10})
        figure.update_xaxes(nticks=4, row=positions[2][0], col=positions[2][1])
        figure.update_xaxes(nticks=4, row=positions[7][0], col=positions[7][1])
    figure.update_xaxes(tickformat=".0%", row=positions[4][0], col=positions[4][1])
    figure.update_xaxes(tickformat=".0%", row=positions[5][0], col=positions[5][1])
    figure.update_yaxes(tickformat=".0%", row=positions[7][0], col=positions[7][1])
    figure.update_xaxes(title_text="\u00cdndice original", row=positions[6][0], col=positions[6][1])
    figure.update_yaxes(
        title_text="\u00cdndice sin precios cr\u00edticos",
        row=positions[6][0],
        col=positions[6][1],
    )
    figure.update_xaxes(tickformat="%d %b", row=positions[2][0], col=positions[2][1])
    figure.update_xaxes(tickformat="%d %b", row=positions[7][0], col=positions[7][1])
    figure.update_xaxes(
        gridcolor="rgba(158,246,229,.11)",
        linecolor="rgba(158,246,229,.25)",
        zerolinecolor="rgba(158,246,229,.25)",
        tickfont={"color": "#a5aeaa"},
        title_font={"color": "#a5aeaa"},
    )
    figure.update_yaxes(
        gridcolor="rgba(158,246,229,.11)",
        linecolor="rgba(158,246,229,.25)",
        zerolinecolor="rgba(158,246,229,.25)",
        tickfont={"color": "#a5aeaa"},
        title_font={"color": "#a5aeaa"},
    )
    figure.update_annotations(font_color="#edf1ef")
    file_name = "dashboard_mobile.html" if mobile else "dashboard_pricing_canasta.html"
    path = OUTPUT_DIR / file_name
    figure.write_html(
        path,
        include_plotlyjs=True,
        full_html=True,
        config=DASHBOARD_CONFIG,
        default_width="100%",
    )
    notices = (chain_notice, province_notice)
    _enhance_dashboard_html(path, validation, notices)

    site_path = SITE_DIR / ("mobile.html" if mobile else "index.html")
    figure.write_html(
        site_path,
        include_plotlyjs=True,
        full_html=True,
        config=DASHBOARD_CONFIG,
        default_width="100%",
    )
    _enhance_dashboard_html(site_path, validation, notices)
    return path


def _write_executive_summary(frames: dict[str, pd.DataFrame], validation: ValidationResult) -> Path:
    publishable_index = frames["index_daily"]
    publishable_index = publishable_index[publishable_index["publishable"]]
    if publishable_index.empty:
        index_finding = "no hubo un índice común publicable en la ventana."
    else:
        index_latest = publishable_index[
            publishable_index["snapshot_date"] == publishable_index["snapshot_date"].max()
        ].sort_values("price_index")
        low_index, high_index = index_latest.iloc[0], index_latest.iloc[-1]
        index_finding = (
            f"sobre {int(low_index['common_gtins']):,} GTIN comunes, "
            f"{low_index['banner_label']} presentó el índice más bajo "
            f"({low_index['price_index']:.2f}) y {high_index['banner_label']} el más alto "
            f"({high_index['price_index']:.2f}). La diferencia es descriptiva del panel "
            "común, no de todo el surtido."
        )

    dispersion = frames["dispersion_entity"]
    chain_panel = _latest_publishable_level(dispersion, "BANNER")
    province_panel = _latest_publishable_level(dispersion, "PROVINCE")
    combo_panel = _latest_publishable_level(dispersion, "BANNER_PROVINCE")
    chain_disp = chain_panel[chain_panel["median_dispersion_clean"].notna()].sort_values(
        "median_dispersion_clean"
    )
    province_disp = province_panel[
        province_panel["median_dispersion_clean"].notna()
    ].sort_values("median_dispersion_clean")
    combo_disp = combo_panel[combo_panel["median_dispersion_clean"].notna()].sort_values(
        "median_dispersion_clean"
    )
    _, chain_context = _dispersion_panel_text(
        "Dispersión por cadena", chain_panel, validation.as_of_date
    )
    _, province_context = _dispersion_panel_text(
        "Dispersión por provincia", province_panel, validation.as_of_date
    )
    _, combo_context = _dispersion_panel_text(
        "Dispersión cadena-provincia", combo_panel, validation.as_of_date
    )
    if chain_disp.empty:
        chain_finding = chain_context
    else:
        high_chain_disp = chain_disp.iloc[-1]
        chain_finding = (
            f"{chain_context} {high_chain_disp['banner_label']} tuvo la mayor dispersión "
            "mediana limpia "
            f"({high_chain_disp['median_dispersion_clean']:.1%}) sobre "
            f"{int(high_chain_disp['publishable_products']):,} productos publicables."
        )
    if province_disp.empty:
        province_finding = province_context
    else:
        high_province_disp = province_disp.iloc[-1]
        province_finding = (
            f"{province_context} {high_province_disp['provincia_codigo']} registró la mayor "
            "dispersión mediana "
            f"limpia ({high_province_disp['median_dispersion_clean']:.1%})."
        )
    if combo_disp.empty:
        combo_finding = combo_context
    else:
        high_combo_disp = combo_disp.iloc[-1]
        combo_finding = (
            f"{combo_context} La combinación más alta fue {high_combo_disp['banner_label']} / "
            f"{high_combo_disp['provincia_codigo']} "
            f"({high_combo_disp['median_dispersion_clean']:.1%})."
        )

    savings = frames["basket_savings"]
    savings = (
        savings[savings["snapshot_date"] == savings["snapshot_date"].max()]
        if not savings.empty
        else savings
    )
    banner_savings = savings[savings["comparison_level"] == "BANNER"]
    province_savings = savings[savings["comparison_level"] == "PROVINCE"]
    if banner_savings.empty:
        banner_saving_text = "sin comparación publicable entre cadenas"
    else:
        banner_saving = banner_savings.iloc[0]
        banner_saving_text = (
            f"${banner_saving['potential_saving_amount']:,.0f} "
            f"({banner_saving['potential_saving_pct']:.1f}%)"
        )
    if province_savings.empty:
        province_saving_text = "sin comparación publicable entre provincias"
    else:
        province_saving = province_savings.iloc[0]
        province_saving_text = (
            f"${province_saving['potential_saving_amount']:,.0f} "
            f"({province_saving['potential_saving_pct']:.1f}%)"
        )
    sensitivity = frames["index_sensitivity"]
    max_sensitivity = sensitivity["sensitivity_delta"].max()
    sensitivity_text = (
        f"{max_sensitivity:.2f} puntos"
        if pd.notna(max_sensitivity)
        else "sin estimación publicable"
    )
    unhealthy = int((~frames["source_health"]["source_healthy"]).sum())
    day_phrase = "1 día nacional" if unhealthy == 1 else f"{unhealthy} días nacionales"

    content = f"""# Resumen ejecutivo

## Objetivo

Comparar nivel de precios, dispersión y costo de una canasta exacta sin mezclar
catálogos, presentaciones ni códigos locales.

## Hallazgos cuantificados

1. **Nivel de precios:** {index_finding}
2. **Dispersión por cadena:** {chain_finding}
3. **Dispersión geográfica:** {province_finding} {combo_finding}
4. **Ahorro potencial de canasta:** entre cadenas publicables, la brecha es {banner_saving_text}. Entre provincias, {province_saving_text}.
5. **Robustez y fuente:** excluir precios críticos cambió los índices como máximo {sensitivity_text}. Hubo {day_phrase} debajo del umbral de salud; esos cortes se identifican visualmente y no prueban cambios comerciales.

## Acciones de negocio

1. El responsable de pricing debe revisar primero los 15 impulsores de mayor contribución de cada cadena y confirmar si son decisiones comerciales o errores de escala.
2. Gestión de categorías debe usar `basket_candidate_review.csv` para evaluar sustituciones; la canasta publicada no cambia automáticamente.
3. Las provincias y combinaciones con estado `SUPPRESSED` no deben aparecer en clasificaciones; se requiere ampliar sucursales antes de decidir.
4. Para negociar precios, priorizar GTIN con dispersión alta persistente durante varios días y no picos de una sola fecha.
5. Archivar al menos 28 días antes de formular conclusiones estructurales o de tendencia.

## Limitaciones específicas

- El índice mide un panel común con igual peso por producto y cadena; no representa participación de mercado ni gasto del consumidor.
- La dispersión P90-P10 describe precios publicados, no promociones ni precios efectivamente pagados.
- El ahorro de canasta existe solo donde los ocho GTIN están presentes; no se imputan faltantes.
- El dato nacional es la red de sucursales observada y no una estimación ponderada por población.
- Siete días permiten describir el corte, no medir inflación, estacionalidad o superioridad estructural.
"""
    path = OUTPUT_DIR / "resumen_ejecutivo.md"
    path.write_text(content, encoding="utf-8")
    return path


def export_all(database_path: Path = DATABASE_PATH) -> None:
    ensure_directories()
    validation = validate(database_path=database_path, raise_on_failure=True)
    connection = duckdb.connect(str(database_path), read_only=True)
    try:
        frames = _export_tables(connection)
        excel_path = _write_excel(frames, validation)
        dashboard_path = _write_dashboard(frames, validation)
        mobile_path = _write_dashboard(frames, validation, mobile=True)
        summary_path = _write_executive_summary(frames, validation)
        metadata = {
            "generated_at": datetime.now(UTC).isoformat(),
            "build_id": validation.build_id,
            "as_of_date": str(validation.as_of_date),
            "confidence": validation.confidence,
            "excel": excel_path.name,
            "dashboard": dashboard_path.name,
            "mobile_dashboard": mobile_path.name,
            "summary": summary_path.name,
            "sample_exports": {
                "dispersion_product_sample": "maximo 10000 filas",
                "price_anomalies_sample": "maximo 10000 filas; criticas primero",
            },
        }
        (OUTPUT_DIR / "run_metadata.json").write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
        )
    finally:
        connection.close()
